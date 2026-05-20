#!/usr/bin/env python3
"""Generate analysis report from Excel files"""

import sys
import os
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl not available")
    sys.exit(1)

def extract_excel_data(excel_path):
    """Extract data from Excel file"""
    wb = load_workbook(excel_path)

    # Get Summary sheet
    summary = wb['Summary']
    device_info = summary['A1'].value

    # Extract device ID and PCIe bus from header
    # Format: "GPU Step Execution Summary - Device X (PCIe: Y)"
    import re
    device_match = re.search(r'Device (\d+) \(PCIe: ([^)]+)\)', device_info)
    device_id = device_match.group(1) if device_match else "unknown"
    pcie_bus = device_match.group(2) if device_match else "unknown"

    # Extract step data from Summary sheet (starting from row 4)
    steps = []
    for row in range(4, summary.max_row + 1):
        step_name = summary.cell(row=row, column=1).value
        if step_name and step_name not in ['Step Name', 'Step Type', None]:
            steps.append({
                'step_name': step_name,  # Full NVTX text name
                'step_type': 'forward' if 'forward_step' in step_name else ('backward' if 'backward_step' in step_name else 'optimizer'),
                'process_sig': summary.cell(row=row, column=2).value,
                'gpu_start_ns': summary.cell(row=row, column=3).value,
                'gpu_end_ns': summary.cell(row=row, column=4).value,
                'duration_ms': summary.cell(row=row, column=5).value,
                'streams_count': summary.cell(row=row, column=6).value,
            })

    # Extract per-stream data for each step type
    stream_data = {}
    for sheet_name in ['forward_step', 'backward_step', 'optimizer_step']:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            stream_data[sheet_name] = []

            # Parse the sheet - data is organized by step sections
            current_step = None
            current_step_name = None
            current_streams = []

            for row in range(1, ws.max_row + 1):
                cell_val = ws.cell(row=row, column=1).value

                # Check if this is a step header row - now contains full NVTX text
                if cell_val and ('megatron.' in str(cell_val) or 'torch.optim' in str(cell_val)):
                    if current_step and current_streams:
                        stream_data[sheet_name].append({
                            'step_num': current_step,
                            'step_name': current_step_name,
                            'streams': current_streams
                        })
                    current_step = cell_val
                    current_step_name = cell_val
                    current_streams = []
                # Check if this is stream data row (starts with stream ID like '7', '165', etc.)
                elif cell_val and str(cell_val).isdigit():
                    stream_id = str(cell_val)
                    gpu_start = ws.cell(row=row, column=2).value
                    gpu_end = ws.cell(row=row, column=3).value
                    duration = ws.cell(row=row, column=4).value
                    percentage = ws.cell(row=row, column=5).value

                    if gpu_start and duration:
                        current_streams.append({
                            'stream_id': stream_id,
                            'gpu_start_ns': gpu_start,
                            'gpu_end_ns': gpu_end,
                            'duration_ms': duration,
                            'percentage': percentage
                        })

            # Add last step
            if current_step and current_streams:
                stream_data[sheet_name].append({
                    'step_num': current_step,
                    'step_name': current_step_name,
                    'streams': current_streams
                })

    return {
        'device_id': device_id,
        'pcie_bus': pcie_bus,
        'steps': steps,
        'stream_data': stream_data
    }


def generate_markdown_report(excel_files, output_path):
    """Generate Markdown report from Excel files"""

    # Sort files by PCIe bus
    device_data = []
    for excel_file in sorted(excel_files):
        data = extract_excel_data(excel_file)
        device_data.append(data)

    # Sort by PCIe bus for consistent ordering
    device_data.sort(key=lambda x: x['pcie_bus'])

    with open(output_path, 'w') as f:
        f.write("# NSYS Profile Analysis Report\n\n")
        f.write("**Iteration:** 16  \n")
        f.write("**Analysis Date:** 2026-05-19  \n")
        f.write("**Model:** DeepSeek-V3  \n\n")

        f.write("---\n\n")

        # Generate table for each device
        for device in device_data:
            pcie = device['pcie_bus']
            device_id = device['device_id']

            f.write(f"## Device {device_id} (PCIe: {pcie})\n\n")

            # Collect all unique stream IDs across all steps
            all_streams = set()
            for step_type in ['forward_step', 'backward_step', 'optimizer_step']:
                if step_type in device['stream_data']:
                    for step in device['stream_data'][step_type]:
                        for stream in step['streams']:
                            all_streams.add(stream['stream_id'])

            sorted_streams = sorted(all_streams, key=lambda x: int(x) if x.isdigit() else x)

            # Table header
            f.write("| Step Name | Total GPU Duration (ms) | Streams Count |")
            for stream_id in sorted_streams:
                f.write(f" Stream {stream_id} (ms) |")
            f.write("\n")

            # Separator
            f.write("|" + "---|" * (3 + len(sorted_streams)) + "\n")

            # Sort steps by GPU start time and write them
            all_steps = sorted(device['steps'], key=lambda x: x['gpu_start_ns'] if x['gpu_start_ns'] else float('inf'))

            # Write steps in chronological order
            for step in all_steps:
                # Use the full step name from NVTX text
                step_name = step['step_name']
                # Truncate if too long for display
                if len(step_name) > 70:
                    step_name = step_name[:67] + "..."
                f.write(f"| {step_name} | {step['duration_ms']:.2f} | {step['streams_count']} |")

                # Get stream data for this step
                streams_for_step = {}
                step_type = step['step_type'] + '_step'
                if step_type in device['stream_data']:
                    # Find the matching step by step name
                    for s in device['stream_data'][step_type]:
                        if s['step_name'] == step['step_name']:
                            for stream in s['streams']:
                                streams_for_step[stream['stream_id']] = stream['duration_ms']
                            break

                for stream_id in sorted_streams:
                    duration = streams_for_step.get(stream_id, '-')
                    if isinstance(duration, (int, float)):
                        f.write(f" {duration:.2f} |")
                    else:
                        f.write(f" {duration} |")
                f.write("\n")

            f.write("\n")

        # Summary section
        f.write("---\n\n")
        f.write("## Summary\n\n")
        f.write("### GPU Devices\n\n")
        f.write("| Device ID | PCIe Bus | GPU Model |\n")
        f.write("|-----------|----------|-----------|\n")
        for device in device_data:
            f.write(f"| {device['device_id']} | {device['pcie_bus']} | NVIDIA GB200 |\n")

        f.write("\n")

        # Calculate totals
        total_forward_time = 0
        total_backward_time = 0
        total_optimizer_time = 0
        forward_count = 0
        backward_count = 0
        optimizer_count = 0

        for device in device_data:
            for step in device['steps']:
                if step['step_type'] == 'forward':
                    total_forward_time += step['duration_ms']
                    forward_count += 1
                elif step['step_type'] == 'backward':
                    total_backward_time += step['duration_ms']
                    backward_count += 1
                elif step['step_type'] == 'optimizer':
                    total_optimizer_time += step['duration_ms']
                    optimizer_count += 1

        f.write("### Average Step Times (across all devices)\n\n")
        if forward_count > 0:
            f.write(f"- **Forward Step:** {total_forward_time / forward_count:.2f} ms (avg across {forward_count} steps)\n")
        if backward_count > 0:
            f.write(f"- **Backward Step:** {total_backward_time / backward_count:.2f} ms (avg across {backward_count} steps)\n")
        if optimizer_count > 0:
            f.write(f"- **Optimizer Step:** {total_optimizer_time / optimizer_count:.2f} ms (avg across {optimizer_count} steps)\n")

        f.write("\n")

        f.write("### Output Files\n\n")
        f.write("Generated files in `documents/train_analyse/example/`:\n\n")
        f.write("**Excel Files (per device):**\n")
        for device in device_data:
            pcie_safe = device['pcie_bus'].replace(':', '_')
            f.write(f"- `iteration_16_gpu_times_{pcie_safe}.xlsx`\n")

        f.write("\n**JSON Files (gpu0 only):**\n")
        f.write("- `iteration_16_forward_step_0008_01_00.0_v3.json`\n")
        f.write("- `iteration_16_backward_step_0008_01_00.0_v3.json`\n")
        f.write("- `iteration_16_optimizer_step_0008_01_00.0_v3.json`\n")
        f.write("- Stream trees: `iteration_16_*_stream_*_0008_01_00.0_v3.json`\n")

    print(f"Report generated: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python generate_report.py <example_dir>")
        sys.exit(1)

    example_dir = Path(sys.argv[1])
    excel_files = list(example_dir.glob("iteration_16_gpu_times_*.xlsx"))

    if not excel_files:
        print(f"No Excel files found in {example_dir}")
        sys.exit(1)

    output_path = example_dir / "analysis_report.md"
    generate_markdown_report(excel_files, output_path)
