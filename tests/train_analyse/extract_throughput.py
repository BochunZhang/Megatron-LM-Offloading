#!/usr/bin/env python3
"""
Extract throughput per GPU from train.log files and export to Excel.

Usage:
    python extract_throughput.py <train_log_path> <output_xlsx_path>

Example:
    python extract_throughput.py /path/to/train.log /path/to/output.xlsx
"""

import argparse
import os
import re

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def extract_data_from_log(log_path: str) -> list:
    """
    Extract iteration and throughput data from train.log file.

    Returns:
        list of dicts with keys: iteration, throughput
    """
    results = []

    try:
        with open(log_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
    except Exception as e:
        print(f"  Warning: Failed to read {log_path}: {e}")
        return results

    # Look for pattern: "iteration XX/ YY | ... | throughput per GPU (TFLOP/s/GPU): ZZZZ"
    # Match the whole line and extract iteration number and throughput value
    pattern = r'iteration\s+(\d+)/\s*\d+.*?throughput per GPU \(TFLOPS?/s/GPU\):\s*([\d.]+)'
    matches = re.findall(pattern, content, re.IGNORECASE)

    for match in matches:
        iteration = int(match[0])
        throughput = float(match[1])
        results.append({
            'iteration': iteration,
            'throughput': throughput
        })

    return results


def generate_excel(output_path: str, data: list):
    """Generate Excel file with iteration and throughput data."""
    if not OPENPYXL_AVAILABLE:
        print("Error: openpyxl not available. Install with: pip install openpyxl")
        return False

    # Ensure output directory exists
    output_dir = os.path.dirname(os.path.abspath(output_path))
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Throughput"

    # Header style
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    center_align = Alignment(horizontal='center', vertical='center')

    # Headers: iteration & Throughput(TFlops/s/GPU)
    headers = ['iteration', 'Throughput(TFlops/s/GPU)']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Data rows
    for row_idx, item in enumerate(data, 2):
        ws.cell(row=row_idx, column=1, value=item['iteration'])
        ws.cell(row=row_idx, column=2, value=item['throughput'])

        # Center alignment for data cells
        for col in range(1, 3):
            ws.cell(row=row_idx, column=col).alignment = center_align

    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 25

    # Save workbook
    wb.save(output_path)
    print(f"  Throughput data saved to: {output_path}")
    return True


def main():
    parser = argparse.ArgumentParser(
        description='Extract throughput per GPU from train.log and export to Excel'
    )
    parser.add_argument('train_log', help='Input train.log file path')
    parser.add_argument('output_xlsx', help='Output Excel file path')

    args = parser.parse_args()

    if not OPENPYXL_AVAILABLE:
        print("Error: openpyxl not available. Install with: pip install openpyxl")
        return 1

    train_log = os.path.abspath(args.train_log)
    if not os.path.exists(train_log):
        print(f"Error: File not found: {train_log}")
        return 1

    print(f"Processing: {train_log}")
    data = extract_data_from_log(train_log)

    if data:
        print(f"  Found {len(data)} iteration records")
        generate_excel(args.output_xlsx, data)
    else:
        print(f"  No throughput data found")
        return 1

    return 0


if __name__ == '__main__':
    exit(main())
